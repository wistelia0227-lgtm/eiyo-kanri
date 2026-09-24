# -*- coding: utf-8 -*-
"""日本食品標準成分表（八訂）増補2023年 → アプリが読む JavaScript のデータ

出すもの:
  js/foods_data.js   本表 54 項目 ＋ 脂肪酸成分表の総量 5 項目 = 59（いつも読む）
  js/foods_amino.js  アミノ酸成分表 第1表（必要になった時だけ読む）
  js/foods_fat.js    脂肪酸成分表 第1表の個別脂肪酸（同上）
  js/foods_carb.js   炭水化物成分表 本表の糖類の内訳（同上）

なぜ分けるか: 本表だけで約 1MB ある。アミノ酸・個別脂肪酸・糖類は、使う場面が限られる割に大きい。
  いつも読み込むと起動が重くなるので、食品の詳しい画面を開いた時だけ読む形にした。

記号の処理（本表の凡例のとおり）:
  数値    -> 数値                    flag '.'
  (数値)  -> 数値（推定値・計算値）   flag 'e'
  Tr      -> 0（微量）               flag 't'
  (Tr)    -> 0（推定で微量）          flag 'T'
  (0)     -> 0（推定ゼロ）           flag 'e'
  -       -> null（未測定）          flag '-'
  空欄    -> null                    flag '-'
  数値†   -> 数値（脚注つき）         flag 'd'
  *       -> null（本表に値が無い。備考の「第3章参照」）flag '*'
flags は食品ごとの 1 文字/列の文字列。全列 '.' の食品は "" にして省く。

使い方: python tools/convert_foods.py
  元の xlsx は ../eiyo-kanri_research/data/ に置いてある（文部科学省から取得したもの）。
"""
import json
import re
import os
import sys
import collections
import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
APP = os.path.dirname(HERE)
SRC = os.path.join(os.path.dirname(APP), 'eiyo-kanri_research', 'data')
OUT = os.path.join(APP, 'js')

MAIN = os.path.join(SRC, '20260327-mxt_kagsei-mext-000029402_02.xlsx')   # 本表
AMINO = os.path.join(SRC, '20260327-mxt_kagsei-mext-000029402_04.xlsx')  # アミノ酸成分表 第1表
FAT = os.path.join(SRC, '20260327-mxt_kagsei-mext-000029402_09.xlsx')    # 脂肪酸成分表 第1表
CARB = os.path.join(SRC, '20260327-mxt_kagsei-mext-000029402_13.xlsx')   # 炭水化物成分表 本表

# (出力キー, 成分識別子, 表示名, 単位)
MAIN_COLS = [
    ('refuse', 'REFUSE', '廃棄率', '%'),
    ('kcal', 'ENERC_KCAL', 'エネルギー', 'kcal'),
    ('kj', 'ENERC', 'エネルギー', 'kJ'),
    ('water', 'WATER', '水分', 'g'),
    ('protcaa', 'PROTCAA', 'アミノ酸組成によるたんぱく質', 'g'),
    ('prot', 'PROT-', 'たんぱく質', 'g'),
    ('fatnlea', 'FATNLEA', '脂肪酸のトリアシルグリセロール当量', 'g'),
    ('chole', 'CHOLE', 'コレステロール', 'mg'),
    ('fat', 'FAT-', '脂質', 'g'),
    ('choavlm', 'CHOAVLM', '利用可能炭水化物(単糖当量)', 'g'),
    ('choavl', 'CHOAVL', '利用可能炭水化物(質量計)', 'g'),
    ('choavldf', 'CHOAVLDF-', '差引き法による利用可能炭水化物', 'g'),
    ('fib', 'FIB-', '食物繊維総量', 'g'),
    ('polyl', 'POLYL', '糖アルコール', 'g'),
    ('cho', 'CHOCDF-', '炭水化物', 'g'),
    ('oa', 'OA', '有機酸', 'g'),
    ('ash', 'ASH', '灰分', 'g'),
    ('na', 'NA', 'ナトリウム', 'mg'),
    ('k', 'K', 'カリウム', 'mg'),
    ('ca', 'CA', 'カルシウム', 'mg'),
    ('mg', 'MG', 'マグネシウム', 'mg'),
    ('p', 'P', 'リン', 'mg'),
    ('fe', 'FE', '鉄', 'mg'),
    ('zn', 'ZN', '亜鉛', 'mg'),
    ('cu', 'CU', '銅', 'mg'),
    ('mn', 'MN', 'マンガン', 'mg'),
    ('iod', 'ID', 'ヨウ素', 'μg'),
    ('se', 'SE', 'セレン', 'μg'),
    ('cr', 'CR', 'クロム', 'μg'),
    ('mo', 'MO', 'モリブデン', 'μg'),
    ('retol', 'RETOL', 'レチノール', 'μg'),
    ('carta', 'CARTA', 'α-カロテン', 'μg'),
    ('cartb', 'CARTB', 'β-カロテン', 'μg'),
    ('crypxb', 'CRYPXB', 'β-クリプトキサンチン', 'μg'),
    ('cartbeq', 'CARTBEQ', 'β-カロテン当量', 'μg'),
    ('vita', 'VITA_RAE', 'ビタミンA(レチノール活性当量)', 'μg'),
    ('vitd', 'VITD', 'ビタミンD', 'μg'),
    ('vite', 'TOCPHA', 'ビタミンE(α-トコフェロール)', 'mg'),
    ('tocphb', 'TOCPHB', 'β-トコフェロール', 'mg'),
    ('tocphg', 'TOCPHG', 'γ-トコフェロール', 'mg'),
    ('tocphd', 'TOCPHD', 'δ-トコフェロール', 'mg'),
    ('vitk', 'VITK', 'ビタミンK', 'μg'),
    ('b1', 'THIA', 'ビタミンB1', 'mg'),
    ('b2', 'RIBF', 'ビタミンB2', 'mg'),
    ('nia', 'NIA', 'ナイアシン', 'mg'),
    ('ne', 'NE', 'ナイアシン当量', 'mg'),
    ('b6', 'VITB6A', 'ビタミンB6', 'mg'),
    ('b12', 'VITB12', 'ビタミンB12', 'μg'),
    ('fol', 'FOL', '葉酸', 'μg'),
    ('pantac', 'PANTAC', 'パントテン酸', 'mg'),
    ('biot', 'BIOT', 'ビオチン', 'μg'),
    ('vitc', 'VITC', 'ビタミンC', 'mg'),
    ('alc', 'ALC', 'アルコール', 'g'),
    ('nacl', 'NACL_EQ', '食塩相当量', 'g'),
]
FAT_COLS = [
    ('fasat', 'FASAT', '飽和脂肪酸', 'g'),
    ('fams', 'FAMS', '一価不飽和脂肪酸', 'g'),
    ('fapu', 'FAPU', '多価不飽和脂肪酸', 'g'),
    ('n3', 'FAPUN3', 'n-3系多価不飽和脂肪酸', 'g'),
    ('n6', 'FAPUN6', 'n-6系多価不飽和脂肪酸', 'g'),
]

stats = collections.Counter()
oddities = []


def parse(v, ctx=''):
    """-> (value, flag)"""
    if v is None:
        stats['blank'] += 1
        return None, '-'
    if isinstance(v, (int, float)):
        stats['num'] += 1
        return clean(v), '.'
    s = str(v).strip().replace('（', '(').replace('）', ')')
    if s == '':
        stats['blank'] += 1
        return None, '-'
    if s in ('-', '－', '―', '‐'):
        stats['-'] += 1
        return None, '-'
    if s == '*':
        # ヨウ素の 3 件（れんこん甘酢・人乳・ぽん酢しょうゆ）。備考に「第3章参照」とある
        stats['*'] += 1
        return None, '*'
    if s == 'Tr':
        stats['Tr'] += 1
        return 0, 't'
    if s == '(Tr)':
        stats['(Tr)'] += 1
        return 0, 'T'
    m = re.fullmatch(r'\(([0-9.]+)\)', s)
    if m:
        stats['(N)'] += 1
        return clean(float(m.group(1))), 'e'
    m = re.fullmatch(r'([0-9.]+)†', s)
    if m:
        stats['N†'] += 1
        oddities.append((ctx, s))
        return clean(float(m.group(1))), 'd'
    if re.fullmatch(r'[0-9.]+', s):
        stats['num'] += 1
        return clean(float(s)), '.'
    stats['UNKNOWN'] += 1
    oddities.append((ctx, s))
    return None, '-'


def clean(x):
    x = round(float(x), 4)
    return int(x) if x == int(x) else x


def header_map(ws, id_row):
    for row in ws.iter_rows(min_row=id_row, max_row=id_row, values_only=True):
        return {str(v).strip(): j for j, v in enumerate(row) if v is not None}
    return {}


def find_id_row(ws, limit=16):
    """成分識別子の行を探す（別冊は 5 行目、本表は 12 行目）"""
    for i, row in enumerate(ws.iter_rows(min_row=1, max_row=limit, values_only=True)):
        vals = [str(v).strip() for v in row if v is not None]
        if '成分識別子' in vals:
            return i + 1
    return None


def read_side(path, skip_ids):
    """別冊を読む。-> (列の定義, {食品番号: [(値, flag), ...]}, 更新日)
    skip_ids = 本表と同じで持つ必要のない識別子"""
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb.worksheets[0]
    idr = find_id_row(ws)
    hm = header_map(ws, idr)
    # 見出しの日本語名は識別子の 2 行上（項目名が複数行に割れているので、あるものを拾う）
    names = {}
    rows = list(ws.iter_rows(min_row=1, max_row=idr, values_only=True))
    for r in rows[:idr - 1]:
        for j, v in enumerate(r):
            if v is None:
                continue
            t = str(v).strip().replace('\n', '')
            if t and j not in names:
                names[j] = t
            elif t and names.get(j) and t not in names[j]:
                names[j] = names[j] + t
    units = {}
    for row in ws.iter_rows(min_row=idr + 1, max_row=idr + 1, values_only=True):
        for j, v in enumerate(row):
            if v:
                units[j] = str(v).strip().replace('/100 g', '').strip()
    cols = []
    for ident, j in hm.items():
        if ident in ('成分識別子',) or ident in skip_ids:
            continue
        cols.append((ident.strip().lower().replace('-', '_'), ident, names.get(j, ident), units.get(j, '')))
    cols.sort(key=lambda c: hm[c[1]])
    data = {}
    for row in ws.iter_rows(min_row=idr + 2, values_only=True):
        if not row[1]:
            continue
        no = str(row[1]).strip().zfill(5)
        data[no] = [parse(row[hm[c[1]]], no + ':' + c[1]) for c in cols]
    wb.close()
    return cols, data


def write_side(filename, varname, title, srcfile, cols, data):
    foods = []
    for no in sorted(data):
        vals, flags = [], []
        for v, f in data[no]:
            vals.append(v)
            flags.append(f)
        fl = ''.join(flags)
        if set(fl) == {'.'}:
            fl = ''
        foods.append([no, fl] + vals)
    out = {
        'meta': {
            'title': title,
            'citation': '日本食品標準成分表（八訂）増補2023年から引用',
            'source_file': os.path.basename(srcfile),
            'basis': '可食部100g当たり',
            'head': ['no', 'flags'],
            'nutrients': [{'key': k, 'id': i, 'name': n, 'unit': u} for k, i, n, u in cols],
            'count': len(foods),
        },
        'foods': foods,
    }
    js = json.dumps(out, ensure_ascii=False, separators=(',', ':'))
    p = os.path.join(OUT, filename)
    with open(p, 'w', encoding='utf-8', newline='\n') as f:
        f.write('/* 出典: 日本食品標準成分表（八訂）増補2023年（文部科学省）から引用 */\nwindow.%s=%s;\n' % (varname, js))
    print('%-22s %5d 食品 × %2d 項目  %6.0f KB' % (filename, len(foods), len(cols), os.path.getsize(p) / 1024))


def main():
    wb = openpyxl.load_workbook(MAIN, read_only=True, data_only=True)
    ws = wb.worksheets[0]
    updated = None
    for row in ws.iter_rows(min_row=1, max_row=1, values_only=True):
        for v in row:
            if v and '更新日' in str(v):
                updated = str(v).replace('更新日：', '')
    hm = header_map(ws, 12)
    missing = [c[1] for c in MAIN_COLS if c[1] not in hm]
    if missing:
        sys.exit('識別子が見つからない: %s' % missing)
    idx_mark_m = hm['CHOAVLM'] + 1    # '*' = エネルギー計算に単糖当量を使用
    idx_mark_d = hm['CHOAVLDF-'] + 1  # '*' = エネルギー計算に差引き法を使用
    idx_remark = 61

    groups = {}
    for w in wb.worksheets[1:]:
        m = re.match(r'(\d+)(.+)', w.title.strip())
        groups['%02d' % int(m.group(1))] = m.group(2).strip()

    # 脂肪酸成分表 第1表（総量 5 つだけ本表に足す）
    wbf = openpyxl.load_workbook(FAT, read_only=True, data_only=True)
    wsf = wbf.worksheets[0]
    hf = header_map(wsf, find_id_row(wsf))
    fat = {}
    for row in wsf.iter_rows(min_row=7, values_only=True):
        if not row[1]:
            continue
        no = str(row[1]).strip().zfill(5)
        fat[no] = [parse(row[hf[c[1]]], no + ':' + c[1]) for c in FAT_COLS]
    wbf.close()

    foods = []
    nofat = 0
    for row in ws.iter_rows(min_row=13, values_only=True):
        if not row[1]:
            continue
        no = str(row[1]).strip().zfill(5)
        name = str(row[3]).strip()
        vals, flags = [], []
        for key, ident, _, _ in MAIN_COLS:
            v, f = parse(row[hm[ident]], no + ':' + ident)
            vals.append(v)
            flags.append(f)
        if no in fat:
            for v, f in fat[no]:
                vals.append(v)
                flags.append(f)
        else:
            nofat += 1
            for _ in FAT_COLS:
                vals.append(None)
                flags.append('-')
        mark = 'M' if str(row[idx_mark_m] or '').strip() == '*' else ('D' if str(row[idx_mark_d] or '').strip() == '*' else '')
        remark = str(row[idx_remark] or '').strip().replace('\n', ' / ')
        fl = ''.join(flags)
        if set(fl) == {'.'}:
            fl = ''
        foods.append([no, str(row[0]).strip().zfill(2), str(row[2]).strip(), name, mark, fl, remark] + vals)

    allcols = MAIN_COLS + FAT_COLS
    out = {
        'meta': {
            'title': '日本食品標準成分表（八訂）増補2023年',
            'source': '文部科学省 科学技術・学術審議会 資源調査分科会報告',
            'citation': '日本食品標準成分表（八訂）増補2023年から引用',
            'url': 'https://www.mext.go.jp/a_menu/syokuhinseibun/mext_00001.html',
            'source_files': [os.path.basename(MAIN), os.path.basename(FAT)],
            'source_updated': updated,
            'basis': '可食部100g当たり',
            'head': ['no', 'group', 'index', 'name', 'choMark', 'flags', 'remark'],
            'head_note': {
                'choMark': "エネルギー計算に使われた利用可能炭水化物。M=単糖当量, D=差引き法",
                'flags': "nutrients と同じ並びの 1 文字/列。'.'=実測等 'e'=()付き推定・計算値 't'=Tr(値0) 'T'=(Tr)(値0) '-'=未測定(値null) 'd'=†脚注つき '*'=本表に値が無い(備考の第3章参照、値null)。空文字は全列 '.'",
            },
            'nutrients': [{'key': k, 'id': i, 'name': n, 'unit': u} for k, i, n, u in allcols],
            'count': len(foods),
        },
        'groups': groups,
        'foods': foods,
    }
    js = json.dumps(out, ensure_ascii=False, separators=(',', ':'))
    p = os.path.join(OUT, 'foods_data.js')
    with open(p, 'w', encoding='utf-8', newline='\n') as f:
        f.write('/* 出典: 日本食品標準成分表（八訂）増補2023年（文部科学省）から引用 */\nwindow.FOODS_DATA=' + js + ';\n')
    print('%-22s %5d 食品 × %2d 項目  %6.0f KB' % ('foods_data.js', len(foods), len(allcols), os.path.getsize(p) / 1024))
    print('  脂肪酸の行が無い食品:', nofat, '／ 元データの更新日:', updated)
    wb.close()

    # ---- 別冊（必要になった時だけ読む）----
    # 本表と重なる列（水分・たんぱく質・脂質など）は持たない
    skip = {'WATER', 'PROTCAA', 'PROT-', 'FATNLEA', 'FAT-', 'CHOAVLM', 'CHOAVL'}
    cols, data = read_side(AMINO, skip)
    write_side('foods_amino.js', 'FOODS_AMINO', 'アミノ酸成分表編 第1表', AMINO, cols, data)
    cols, data = read_side(FAT, skip | {c[1] for c in FAT_COLS})
    write_side('foods_fat.js', 'FOODS_FAT', '脂肪酸成分表編 第1表', FAT, cols, data)
    cols, data = read_side(CARB, skip)
    write_side('foods_carb.js', 'FOODS_CARB', '炭水化物成分表編 本表', CARB, cols, data)

    print('\n記号の内わけ:', dict(stats))
    if oddities:
        print('変わった書き方（先頭 10 件）:', oddities[:10])


if __name__ == '__main__':
    main()
